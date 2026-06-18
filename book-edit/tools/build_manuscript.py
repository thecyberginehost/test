#!/usr/bin/env python3
"""
Compile the edited chapters into a single standard-format manuscript .docx.

Source : book-edit/02-edited/*.md  (one file per chapter, in filename order)
Output : manuscript_build/Project_ForgePulse_Shadows_of_War.docx  (gitignored)

Format: Shunn-style — Times New Roman 12pt, double-spaced, 1" margins,
0.5" first-line indents, title page (contact + word count), running header
with page numbers, each chapter starting on a fresh page. Chapter
date/time/location stamps are preserved as italic openers.

Requires: python-docx  (pip install python-docx)
Run:      python3 book-edit/tools/build_manuscript.py
"""
import os, glob, re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, "book-edit", "02-edited")
OUT_DIR = os.path.join(ROOT, "manuscript_build")
OUT = os.path.join(OUT_DIR, "Project_ForgePulse_Shadows_of_War.docx")

TITLE, SUB, AUTHOR = "PROJECT FORGEPULSE", "Shadows of War", "D.K. Shadow"
REALNAME, EMAIL = "Anthony Amore", "amorework2@gmail.com"

files = sorted(glob.glob(os.path.join(SRC, "*.md")))
wc = sum(len(open(f, encoding="utf-8").read().split()) for f in files)
wordcount = int(round(wc, -2))

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, m, Inches(1))


def para(text="", align=None, indent=False, italic=False, bold=False, size=None, space_before=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Inches(0.5) if indent else Inches(0)
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if text:
        r = p.add_run(text); r.italic = italic; r.bold = bold
        if size:
            r.font.size = Pt(size)
    return p


# running header: Surname / TITLE / page#
hdr = sec.header.paragraphs[0]; hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hdr.add_run("Shadow / PROJECT FORGEPULSE / ")
r = hdr.add_run()
fc1 = OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"), "begin")
it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "end")
r._r.append(fc1); r._r.append(it); r._r.append(fc2)

# title page
c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.LEFT
c.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
c.add_run(f"{REALNAME}\n{EMAIL}")
w = doc.add_paragraph(); w.alignment = WD_ALIGN_PARAGRAPH.RIGHT
w.add_run(f"approx. {wordcount:,} words")
for _ in range(8):
    doc.add_paragraph()
para(TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=26)
para(SUB, align=WD_ALIGN_PARAGRAPH.CENTER, size=15)
doc.add_paragraph()
para(f"a novel by {AUTHOR}", align=WD_ALIGN_PARAGRAPH.CENTER, size=13)


def inline_runs(p, text):
    """Render **bold** segments; rest plain."""
    i = 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        if m.start() > i:
            p.add_run(text[i:m.start()])
        p.add_run(m.group(1)).bold = True
        i = m.end()
    if i < len(text):
        p.add_run(text[i:])


def _is_meta(ln):
    """Skip editorial notes that live in chapter files but aren't story text."""
    s = ln.strip()
    return s.startswith("*(Provisional") or "DISCREPANCY-REPORT" in s


for f in files:
    ne = [ln.rstrip() for ln in open(f, encoding="utf-8") if ln.strip() and not _is_meta(ln)]
    if not ne:
        continue
    doc.add_page_break()
    head = re.sub(r"^#+\s*", "", ne[0]).replace("*", "").strip()
    para(head, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_before=24)
    doc.add_paragraph()
    first = False
    for ln in ne[1:]:
        t = ln.strip(); low = t.lower()
        if low.startswith(("date:", "time:", "location:")):
            para(t, italic=True); continue
        if t in ("#", "***", "* * *", "---", "* * * *", "___"):
            para("#", align=WD_ALIGN_PARAGRAPH.CENTER); continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.5) if first else Inches(0)
        first = True
        inline_runs(p, re.sub(r"^#+\s*", "", t))

os.makedirs(OUT_DIR, exist_ok=True)
doc.save(OUT)
print(f"SAVED: {OUT} | ~{wordcount:,} words | {len(files)} files")
